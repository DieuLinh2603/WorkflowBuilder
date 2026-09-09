from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from datetime import date


OUTPUT = "Bao-cao-tong-quan-he-thong-WorkflowBuilder.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run.font.name = "Arial"
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=(255, 255, 255))
        shade(table.rows[0].cells[index], "1F4E78")
        if widths:
            table.rows[0].cells[index].width = Cm(widths[index])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(row):
            set_cell_text(cells[col_index], value)
            if row_index % 2:
                shade(cells[col_index], "EAF2F8")
            if widths:
                cells[col_index].width = Cm(widths[col_index])
    doc.add_paragraph()
    return table


def bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.add_run(text)
    return paragraph


def numbered(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.add_run(text)
    return paragraph


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.15
for style_name, size, color in [
    ("Title", 24, (31, 78, 120)),
    ("Heading 1", 17, (31, 78, 120)),
    ("Heading 2", 13, (46, 116, 181)),
    ("Heading 3", 11, (55, 55, 55)),
]:
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(*color)
    style.font.bold = True

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(80)
run = title.add_run("BÁO CÁO TỔNG QUAN HỆ THỐNG\nWORKFLOW BUILDER")
run.bold = True
run.font.name = "Arial"
run.font.size = Pt(24)
run.font.color.rgb = RGBColor(31, 78, 120)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run("Tech stack – Mapping yêu cầu – Khó khăn triển khai").italic = True
doc.add_paragraph()
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run(f"Ngày lập báo cáo: {date.today().strftime('%d/%m/%Y')}")

doc.add_page_break()

doc.add_heading("1. Overview về hệ thống", level=1)
doc.add_heading("1.1. Kết luận tổng quan", level=2)
doc.add_paragraph(
    "Hệ thống hiện tại là một nền tảng Workflow Builder nội bộ kết hợp Data Pipeline. "
    "Hệ thống cho phép thiết kế quy trình trực quan, cấu hình form và điều kiện, publish/version hóa, "
    "chạy workflow, theo dõi task, đồng thời nhận dữ liệu từ CSV, REST API hoặc PostgreSQL để tạo Dataset "
    "và khởi chạy batch workflow."
)
doc.add_paragraph(
    "Mức độ hoàn thiện phù hợp với MVP chức năng tương đối sâu hoặc bản pilot nội bộ. Phần lõi workflow "
    "đã khá đầy đủ, nhưng chưa nên xem là production-ready ở quy mô lớn vì còn thiếu E2E test, CI/CD, "
    "observability, object storage, job queue và một số cơ chế bảo mật/vận hành."
)
note = doc.add_paragraph()
note.add_run("Phạm vi đánh giá: ").bold = True
note.add_run(
    "Repository không có business requirement document hoàn chỉnh; README gốc gần như trống. "
    "Bảng yêu cầu trong báo cáo được tái dựng từ source code, API, database migration, UI và test hiện có."
)

doc.add_heading("1.2. Nhóm người dùng", level=2)
add_table(doc, ["Vai trò", "Trách nhiệm chính"], [
    ("ADMIN", "Quản lý user, group, settings, connector và có quyền quản trị toàn hệ thống."),
    ("WORKFLOW_OWNER", "Tạo workflow, quản lý editor, publish, quản lý pipeline và data binding."),
    ("EDITOR", "Chỉnh sửa các workflow được phân công."),
    ("VIEWER", "Gửi request, theo dõi request và xử lý task nếu được giao."),
], [4.5, 11.5])
doc.add_paragraph(
    "Ngoài role cấp hệ thống, quyền còn được kiểm tra theo owner, editor được gán trực tiếp, audience được phép "
    "gửi request và user/group/role được gán xử lý tại từng step."
)

doc.add_heading("1.3. Kiến trúc tổng thể", level=2)
architecture = doc.add_paragraph()
architecture.alignment = WD_ALIGN_PARAGRAPH.CENTER
architecture.paragraph_format.space_before = Pt(6)
architecture.paragraph_format.space_after = Pt(12)
arch_run = architecture.add_run(
    "React SPA\n↓ REST API + JWT\nSpring Boot Controllers & Services\n"
    "↓\nWorkflow Engine  |  Notification/Reminder  |  Data Pipeline Engine\n"
    "↓\nSpring Data JPA + PostgreSQL + Flyway\n\n"
    "Nguồn pipeline: CSV  |  REST API  |  PostgreSQL"
)
arch_run.font.name = "Consolas"
arch_run.font.size = Pt(10)

doc.add_heading("1.4. Luồng nghiệp vụ chính", level=2)
for item in [
    "Owner tạo workflow draft.",
    "Owner hoặc editor cấu hình step, form và connection.",
    "Hệ thống validate graph trước khi publish.",
    "Workflow được publish thành một version bất biến.",
    "User thuộc audience gửi request đơn hoặc batch.",
    "Runtime engine đi qua graph và tạo task cho người xử lý.",
    "Approval, review hoặc assignment hoàn thành sẽ kích hoạt nhánh tiếp theo.",
    "Notification, system action và end step được xử lý tự động.",
    "Tiến trình được lưu thành instance, task và step log.",
    "Data pipeline có thể tạo Dataset và tự động đẩy record thay đổi vào batch workflow.",
]:
    numbered(doc, item)

doc.add_heading("1.5. Các module hiện có", level=2)
modules = {
    "Authentication và user management": [
        "Đăng nhập email/password bằng JWT stateless và BCrypt.",
        "Đổi mật khẩu; quên/reset mật khẩu bằng token có thời hạn.",
        "Quản lý user, manager, chức danh, nhiều system role và trạng thái active.",
        "Tạo group và quản lý thành viên.",
    ],
    "Workflow Designer": [
        "Canvas React Flow với các node START, APPROVAL, REVIEW, ASSIGNMENT, NOTIFICATION, SYSTEM_ACTION và END.",
        "Thêm/xóa node, drag/drop, lưu layout, nối node và cấu hình connection.",
        "IF/ELSE, AND/OR, priority, condition theo kiểu dữ liệu và biểu thức tính toán.",
        "Kiểm tra graph, connection, cycle và cấu hình trước khi publish.",
    ],
    "Workflow lifecycle và versioning": [
        "DRAFT, PUBLISHED, SUSPENDED, ARCHIVED và DELETED.",
        "Published version bất biến; chỉnh sửa thông qua working copy/version mới.",
        "Compare, restore, duplicate, suspend, archive, chuyển owner và gán editor.",
    ],
    "Dynamic form và runtime": [
        "Field text, number, date, checkbox, file, required, placeholder và display order.",
        "Submission đơn, batch CSV và auto-save draft.",
        "Manual/automatic approval, review, assignment, multi-assignee và dynamic actor.",
        "ANY, ALL, PERCENTAGE; reject quay lại sửa; cancel, withdraw và re-evaluate.",
    ],
    "Notification và reminder": [
        "In-app notification, email, Microsoft Teams webhook và generic webhook.",
        "Deadline reminder scheduler và cảnh báo pipeline thất bại.",
    ],
    "Data Pipeline": [
        "Nguồn CSV, REST API và PostgreSQL SELECT.",
        "REST pagination: page, offset, cursor.",
        "INNER/LEFT/FULL join; cardinality validation; filter, cast, trim, sum và average.",
        "Output schema, composite business key, preview và Dataset versioning.",
        "Manual/once/daily schedule, retry và Data Binding Dataset → batch workflow.",
    ],
}
for module, items in modules.items():
    doc.add_heading(module, level=3)
    for item in items:
        bullet(doc, item)

doc.add_heading("2. Tech stack lựa chọn và lý do", level=1)
doc.add_heading("2.1. Frontend", level=2)
add_table(doc, ["Công nghệ", "Vai trò", "Lý do lựa chọn"], [
    ("React 18", "SPA", "Component hóa tốt; phù hợp màn hình nhiều state như designer, form động và dashboard."),
    ("Vite 5", "Dev server/build", "Khởi động nhanh, cấu hình đơn giản, phù hợp MVP."),
    ("React Router 6", "Routing", "Quản lý public/protected route và route theo role."),
    ("@xyflow/react", "Workflow canvas", "Có sẵn node, edge, handle, viewport, drag/drop; giảm đáng kể chi phí tự xây graph editor."),
    ("Tailwind CSS", "Styling", "Phát triển UI nhanh, dễ chuẩn hóa layout và responsive."),
    ("Lucide React", "Icon", "Bộ icon nhẹ, đồng nhất và dễ sử dụng."),
    ("Fetch wrapper", "API client", "Tự gắn JWT, xử lý 401, lỗi và toast; đủ cho API hiện tại."),
], [3.2, 3.7, 9.2])
doc.add_paragraph(
    "React Flow là lựa chọn phù hợp nhất với phần designer vì bài toán cần xử lý node position, connection handle, "
    "edge rendering, zoom/pan, selection và đồng bộ graph với backend. Nếu tự xây bằng DOM/SVG, chi phí và rủi ro lỗi sẽ cao hơn nhiều."
)

doc.add_heading("Hạn chế frontend", level=3)
for item in [
    "Dùng JavaScript thay vì TypeScript nên payload cấu hình JSON dễ sai type.",
    "Không có ESLint/test script, component test hoặc E2E test.",
    "Bundle JavaScript khoảng 764 KB, chưa có route-level code splitting.",
    "Axios được khai báo nhưng code chính dùng fetch.",
    "Có trang WorkflowBuilderPage cũ chứa mock canvas nhưng route chính không sử dụng.",
    "Dashboard Viewer điều hướng tới /my-work trong khi route thật là /tasks.",
]:
    bullet(doc, item)

doc.add_heading("2.2. Backend", level=2)
add_table(doc, ["Công nghệ", "Vai trò", "Lý do lựa chọn"], [
    ("Java 17", "Ngôn ngữ", "LTS, ổn định và phù hợp hệ thống nghiệp vụ doanh nghiệp."),
    ("Spring Boot 3.2.5", "Backend framework", "Hệ sinh thái đầy đủ cho REST, security, validation, database và scheduler."),
    ("Spring MVC", "REST API", "Phù hợp request/response đồng bộ và quy mô MVP."),
    ("Spring Data JPA", "Persistence", "Giảm boilerplate truy vấn, hỗ trợ transaction."),
    ("Spring Security", "Bảo mật", "JWT filter, role và method-level authorization."),
    ("Jakarta Validation", "Validate DTO", "Chặn request không hợp lệ trước business layer."),
    ("JJWT + BCrypt", "Auth", "JWT stateless và lưu mật khẩu an toàn."),
    ("Jackson", "JSON", "Serialize form snapshot, step config và pipeline definition."),
    ("Flyway", "DB migration", "Schema có version và có thể tái tạo."),
    ("ShedLock", "Scheduler lock", "Ngăn nhiều application instance chạy trùng scheduler job."),
    ("Spring Mail", "Email", "Reset password, notification và reminder."),
    ("Swagger/OpenAPI", "API docs", "Hỗ trợ kiểm thử và trao đổi API."),
    ("JUnit/Mockito", "Unit test", "Kiểm tra workflow rule và service logic."),
    ("Testcontainers", "Integration test", "Kiểm thử Flyway trên PostgreSQL thật."),
], [3.2, 3.7, 9.2])

doc.add_heading("2.3. Database", level=2)
doc.add_paragraph(
    "PostgreSQL phù hợp vì dữ liệu có quan hệ chặt giữa workflow-version-step-connection, user-role-group, "
    "instance-task-history và pipeline-run-dataset-version-record. Publish, task completion và transition cần transaction nhất quán."
)
doc.add_paragraph(
    "Flyway quản lý schema qua các migration có version. H2 được khai báo cho local development, nhưng cấu hình mặc định "
    "vẫn trỏ PostgreSQL; vì vậy H2 chưa phải fallback tự động đúng nghĩa."
)

doc.add_heading("3. Bảng mapping yêu cầu", level=1)
doc.add_paragraph("Quy ước: ✅ Làm được; 🟡 Làm được một phần hoặc chưa đủ production; ❌ Chưa được hỗ trợ trong phiên bản hiện tại.")

requirements = [
    ("Đăng nhập email/password", "✅", "JWT, Spring Security và BCrypt."),
    ("Đổi/quên/reset mật khẩu", "✅", "Token reset có hash, expiry và trạng thái đã sử dụng."),
    ("SSO, LDAP, OAuth2, MFA", "❌", "Chưa có implementation."),
    ("Quản lý user", "✅", "CRUD, activate/deactivate, manager, job title và nhiều role."),
    ("Quản lý group", "✅", "Tạo group, thêm và xóa member."),
    ("RBAC cấp hệ thống", "✅", "Bốn role và có bảo vệ frontend/backend."),
    ("Phân quyền theo workflow", "✅", "Owner, scoped editor và audience."),
    ("Audit thay đổi quyền", "❌", "Chưa có audit log quản trị đầy đủ."),
    ("Tạo workflow", "✅", "Có metadata, type, module và submission mode."),
    ("Thiết kế trên canvas", "✅", "Node/edge, drag/drop và lưu layout."),
    ("Các step nghiệp vụ", "✅", "Đủ 7 loại step chính."),
    ("IF/ELSE và condition", "✅", "Nhiều clause, AND/OR, priority và expression."),
    ("Validate trước publish", "✅", "Kiểm tra start/end, connection, config, field key và cycle."),
    ("Vòng lặp tùy ý", "❌", "Cố ý không hỗ trợ; chỉ cho correction loop từ manual reject/review fail."),
    ("Workflow versioning", "✅", "Working copy, compare, restore và immutable published version."),
    ("Duplicate/suspend/archive", "✅", "Có lifecycle và API tương ứng."),
    ("Form động", "✅", "Text, number, date, checkbox và file."),
    ("Draft form", "✅", "Auto-save cho submission đơn."),
    ("Upload file nghiệp vụ", "🟡", "File encode Data URL và lưu trong JSON snapshot; chưa có object storage."),
    ("Batch CSV", "✅", "Template, parse, validate và parent batch instance."),
    ("Batch có field file", "❌", "Bị cấm; cần dùng URL dạng text."),
    ("Workflow catalog", "✅", "Chỉ hiển thị workflow user được phép submit."),
    ("Approval thủ công/tự động", "✅", "Fixed user, role-based, dynamic actor và condition."),
    ("Multi-approver", "✅", "ANY, ALL và PERCENTAGE."),
    ("Review/trả lại sửa", "✅", "Pass/fail và correction loop."),
    ("Assignment", "✅", "Gán user hoặc group."),
    ("Task inbox/history", "✅", "Pending/completed task, task detail và journey."),
    ("Withdraw/cancel/re-evaluate", "✅", "Có API và business rule."),
    ("Dashboard theo role", "🟡", "Có tổng quan nhưng tính ở client; chưa có analytics backend."),
    ("Route dashboard Viewer", "🟡", "Link /my-work không khớp route /tasks."),
    ("In-app notification", "✅", "Notification center và unread count."),
    ("Email/Teams/webhook", "🟡", "Có code nhưng phụ thuộc cấu hình; chưa có integration test."),
    ("Deadline reminder", "✅", "Scheduler, delivery tracking và setting toàn cục."),
    ("System Action cập nhật dữ liệu", "✅", "Cập nhật field snapshot."),
    ("System Action tạo record", "🟡", "Chỉ tạo record logic trong snapshot, chưa ghi hệ thống ngoài."),
    ("System Action gọi API", "🟡", "Thiếu OAuth, custom headers, response mapping và circuit breaker."),
    ("Connector REST", "🟡", "Chỉ GET; auth none, bearer và API key."),
    ("Connector PostgreSQL", "✅", "Read-only SELECT, timeout và max row."),
    ("MySQL/SQL Server/S3/Excel", "❌", "Chưa hỗ trợ."),
    ("Connector permission", "✅", "Admin tạo và cấp quyền cho user."),
    ("Mã hóa credential", "✅", "Có cipher với random nonce."),
    ("Credential rotation/KMS", "❌", "Chưa có."),
    ("CSV pipeline source", "✅", "Version, checksum và lưu binary."),
    ("REST pagination", "✅", "None, page, offset và cursor."),
    ("Join nhiều nguồn", "✅", "INNER, LEFT, FULL và cardinality validation."),
    ("Filter/cast/trim", "✅", "Có trong guided UI và engine."),
    ("Sum/average", "✅", "Có group by."),
    ("Transform nâng cao", "🟡", "Engine hỗ trợ nhiều hơn UI; cần chỉnh JSON nâng cao."),
    ("Dataset versioning", "✅", "Phân loại NEW, CHANGED và UNCHANGED."),
    ("Phát hiện record bị xóa", "❌", "Chưa có change type DELETED."),
    ("Composite business key", "✅", "Hỗ trợ nhiều cột."),
    ("Pipeline preview", "✅", "Source, join và output preview."),
    ("Lịch pipeline", "🟡", "Chỉ manual, once, daily; timezone cố định Việt Nam."),
    ("Cron/weekly/monthly/event-driven", "❌", "Chưa có."),
    ("Pipeline retry", "✅", "Ba lần với backoff."),
    ("Pipeline bất đồng bộ", "❌", "Manual run chạy đồng bộ trong HTTP request."),
    ("Dataset → Workflow binding", "✅", "Mapping, filter, manual hoặc auto trigger."),
    ("Xóa Data Binding", "❌", "Chưa có DELETE endpoint; có thể inactive bằng update."),
    ("Idempotency pipeline batch", "✅", "Theo dataset version, business key và checksum."),
    ("Swagger/OpenAPI", "✅", "Có API documentation."),
    ("Backend unit test", "✅", "64 test pass trong lần kiểm tra."),
    ("Migration integration test", "🟡", "Có Testcontainers nhưng bị skip do môi trường thiếu Docker."),
    ("Frontend/E2E/load test", "❌", "Chưa có."),
    ("Docker/CI/CD", "❌", "Không thấy cấu hình trong repository."),
    ("Metrics/health/log tập trung", "❌", "Chưa có Actuator/Micrometer/log aggregation."),
    ("Redis/job queue", "❌", "Cache in-memory; chưa có Kafka/RabbitMQ."),
    ("Multi-tenant/mobile/reporting", "❌", "Chưa có tenant model, mobile app hoặc reporting chuyên sâu."),
]
add_table(doc, ["Yêu cầu", "Trạng thái", "Đánh giá"], requirements, [5.2, 2.0, 9.0])

doc.add_paragraph(
    "Các mục đánh dấu ❌ có nghĩa là phiên bản code hiện tại chưa hỗ trợ, không có nghĩa là kiến trúc không thể mở rộng để bổ sung."
)

doc.add_heading("4. Các khó khăn khi thực hiện", level=1)
difficulties = [
    ("4.1. Workflow engine tổng quát",
     "Engine phải xử lý IF/ELSE, approval/reject, review, automatic step, human task, multi-approver, correction loop và transaction failure. "
     "Activation round phải được quản lý để task cũ không bị tính nhầm cho vòng approval mới."),
    ("4.2. Versioning và tính bất biến",
     "Không thể sửa trực tiếp definition đang có instance chạy. Hệ thống phải khóa published version, copy toàn bộ graph sang working copy, "
     "giữ quan hệ family/version và bảo đảm instance cũ chạy theo đúng definition ban đầu."),
    ("4.3. Phân quyền hai tầng",
     "System role chưa đủ; còn phải kiểm tra owner, editor, audience, assignee, requester và batch record recipient. "
     "Frontend guard chỉ phục vụ UX, quyền thật phải nằm ở backend."),
    ("4.4. Batch workflow",
     "Một parent instance chứa nhiều record độc lập giúp giảm số lượng instance/task nhưng khiến runtime, task grouping, summary và update từ pipeline phức tạp hơn."),
    ("4.5. Dữ liệu bên ngoài",
     "API timeout, pagination sai, schema drift, CSV encoding, join key thiếu/trùng và cast lỗi đều cần được kiểm soát. "
     "Engine hiện có guard nhưng vẫn tải toàn bộ dữ liệu vào memory."),
    ("4.6. Scheduler và concurrency",
     "ShedLock tránh chạy trùng nhưng pipeline vẫn chạy tuần tự, external I/O đồng bộ, manual run giữ HTTP request và chưa có worker queue/job cancellation."),
    ("4.7. Quản lý file",
     "File nghiệp vụ đang lưu base64 trong JSON snapshot và CSV pipeline lưu trong PostgreSQL bytea. Cách này phù hợp file nhỏ nhưng làm database, backup và replication nặng; chưa có antivirus scanning hoặc signed URL."),
    ("4.8. Bảo mật integration",
     "Cần outbound allowlist thống nhất, chống SSRF/DNS rebinding, OAuth2, secret manager, rotation, rate limiting, circuit breaker và audit credential. JWT localStorage cũng cần hardening cho production."),
    ("4.9. JSON thiếu type safety",
     "Map<String,Object> và JSON text linh hoạt nhưng khó validate, refactor, sinh OpenAPI schema và giữ contract frontend/backend. Nên chuyển dần sang typed DTO và versioned schema."),
    ("4.10. Test và vận hành",
     "Test hiện tập trung tốt vào service logic nhưng thiếu controller/security integration, external connector, scheduler concurrency, frontend, E2E và load test. "
     "Repo cũng chưa có Docker, CI/CD, metrics, health check, tracing, backup/restore hoặc runbook."),
]
for heading, body in difficulties:
    doc.add_heading(heading, level=2)
    doc.add_paragraph(body)

doc.add_heading("5. Kết quả kiểm tra", level=1)
for item in [
    "Maven backend build thành công.",
    "65 test được phát hiện: 64 pass, 1 Flyway/PostgreSQL integration test bị skip vì môi trường không có Docker.",
    "Frontend đã sinh được bundle dist; esbuild có thể bị spawn EPERM khi chạy trong sandbox hạn chế tiến trình con.",
    "Không có frontend automated test hoặc browser E2E test.",
]:
    bullet(doc, item)

doc.add_heading("6. Đánh giá mức độ sẵn sàng", level=1)
add_table(doc, ["Mức sử dụng", "Đánh giá"], [
    ("Demo nghiệp vụ", "Tốt."),
    ("Pilot nội bộ, ít người dùng", "Có thể triển khai sau khi cấu hình PostgreSQL, SMTP và secret."),
    ("Production nội bộ quy mô vừa", "Cần object storage, monitoring, CI/CD, E2E test và security hardening."),
    ("Dữ liệu lớn/nhiều pipeline đồng thời", "Chưa phù hợp."),
    ("SaaS multi-tenant", "Chưa phù hợp."),
], [6.2, 10.0])

doc.add_heading("7. Ưu tiên trước khi production", level=1)
priorities = [
    "Thêm Docker/CI và test API/E2E.",
    "Chuyển file nghiệp vụ và pipeline file sang object storage.",
    "Chuyển pipeline execution sang job queue/worker.",
    "Thêm Actuator, metrics, centralized logging và alert.",
    "Bổ sung rate limit, refresh/revoke token và outbound URL policy.",
    "Sửa route /my-work thành /tasks.",
    "Chuẩn hóa pipeline payload bằng typed DTO thay cho Map<String,Object>.",
    "Thêm retention policy cho dataset version, pipeline file và audit log.",
    "Viết README kiến trúc, hướng dẫn chạy và business requirement baseline.",
]
for item in priorities:
    numbered(doc, item)

doc.add_heading("8. Kết luận", level=1)
doc.add_paragraph(
    "Phần workflow core đã đáp ứng phần lớn business capability quan trọng: designer, versioning, phân quyền, approval runtime, "
    "batch processing và pipeline binding. Các khoảng trống chính nằm ở enterprise integration, security hardening, scalability, "
    "automated testing và vận hành production. Vì vậy hệ thống phù hợp để demo hoặc pilot nội bộ, nhưng cần hoàn thiện các ưu tiên "
    "trên trước khi triển khai production quy mô vừa hoặc lớn."
)

# Header and footer with page number.
for sec in doc.sections:
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("Workflow Builder – System Review")
    header_run.font.name = "Arial"
    header_run.font.size = Pt(8)
    header_run.font.color.rgb = RGBColor(120, 120, 120)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Trang ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

doc.core_properties.title = "Báo cáo tổng quan hệ thống Workflow Builder"
doc.core_properties.subject = "Tech stack, mapping yêu cầu và khó khăn triển khai"
doc.core_properties.author = "Codex"
doc.save(OUTPUT)
print(OUTPUT)
